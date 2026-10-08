import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".md", ".txt"}


@dataclass
class Chunk:
    id: str
    source: str
    page: int | None
    text: str

    def cite(self) -> str:
        return f"{self.source}, p.{self.page}" if self.page else self.source


def collect_files(inputs: list[Path]) -> list[Path]:
    files: list[Path] = []
    for p in inputs:
        if p.is_dir():
            files += [f for f in sorted(p.rglob("*")) if f.suffix.lower() in SUPPORTED]
        elif p.suffix.lower() in SUPPORTED:
            files.append(p)
        else:
            raise ValueError(f"unsupported input: {p} (supported: {', '.join(sorted(SUPPORTED))})")
    return files


def _read(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return [(i, page.extract_text() or "") for i, page in enumerate(PdfReader(path).pages, 1)]
    if suffix == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        return [(None, "\n\n".join(parts))]
    return [(None, path.read_text(encoding="utf-8", errors="replace"))]


def _split(text: str, size: int, overlap: int) -> list[str]:
    paragraphs = [" ".join(p.split()) for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces, current = [], []
    for para in paragraphs:
        words = para.split()
        if current and len(current) + len(words) > size:
            pieces.append(" ".join(current))
            current = current[-overlap:]
        current.extend(words)
        while len(current) > size * 1.5:
            pieces.append(" ".join(current[:size]))
            current = current[size - overlap:]
    if current:
        pieces.append(" ".join(current))
    return pieces


def chunk_documents(files: list[Path], size: int, overlap: int) -> list[Chunk]:
    chunks: list[Chunk] = []
    for doc_no, path in enumerate(files, 1):
        doc, n = f"D{doc_no}", 0
        for page, text in _read(path):
            for piece in _split(text, size, overlap):
                n += 1
                chunks.append(Chunk(id=f"{doc}-{n:03d}", source=path.name, page=page, text=piece))
    return chunks
